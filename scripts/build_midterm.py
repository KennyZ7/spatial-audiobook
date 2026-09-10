"""Build the midterm DOCX from its reviewed Markdown (requires python-docx/Pillow)."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT=Path(__file__).resolve().parents[1]
FOLDER=ROOT/'docs'/'midterm'

def main():
    # Simple original architecture figure; all labels mirror implemented modules.
    font_path=Path('C:/Windows/Fonts/msyh.ttc')
    if not font_path.exists():
        import os
        font_path=Path(os.environ['REPORT_FONT'])
    im=Image.new('RGB',(1600,410),'white');d=ImageDraw.Draw(im)
    font=ImageFont.truetype(str(font_path),29)
    labels=['原文输入','可编辑声音脚本','配音与素材','实际时间编排','双耳空间渲染','试听与导出']
    for i,label in enumerate(labels):
        x=20+i*264
        d.rounded_rectangle((x,90,x+232,200),radius=12,fill='#e9f0ec',outline='#3e6555',width=2)
        d.text((x+116,145),label,font=font,fill='#253e33',anchor='mm')
        if i<5:
            d.line((x+234,145,x+260,145),fill='#3e6555',width=3)
            d.polygon([(x+260,145),(x+250,139),(x+250,151)],fill='#3e6555')
    d.text((800,275),'工程 JSON 保存原文 角色 事件 时间 轨迹与版本',font=font,fill='#253e33',anchor='mm')
    d.text((800,330),'历史成品快照 → WAV 分轨包 HTML报告 CSV',font=font,fill='#253e33',anchor='mm')
    im.save(FOLDER/'evidence'/'architecture.png')
    doc=Document();sec=doc.sections[0]
    for border in list(doc.styles.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=Cm(2.1);sec.bottom_margin=Cm(2.1);sec.left_margin=Cm(2.3);sec.right_margin=Cm(2.3)
    for name,size,east in [('Normal',11,'宋体'),('Title',20,'黑体'),('Heading 1',15,'黑体'),('Heading 2',12,'黑体')]:
        s=doc.styles[name];s.font.name='Times New Roman';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0)
        s.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),east)
        s.paragraph_format.space_after=Pt(7)
    doc.styles['Normal'].paragraph_format.line_spacing=1.22
    foot=sec.footer.paragraphs[0];foot.alignment=2
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)
    lines=(FOLDER/'中期检查报告.md').read_text('utf-8').splitlines();i=0
    def picture(name,caption):
        doc.add_picture(str(FOLDER/'evidence'/name),width=Cm(16.1))
        p=doc.add_paragraph(caption);p.alignment=1;p.paragraph_format.space_after=Pt(10)
        for r in p.runs:r.font.size=Pt(9)
    while i<len(lines):
        line=lines[i].strip();i+=1
        if not line:continue
        if line.startswith('|'):
            data=[line]
            while i<len(lines) and lines[i].strip().startswith('|'):data.append(lines[i].strip());i+=1
            parsed=[[c.strip() for c in row.strip('|').split('|')] for row in data if not re.fullmatch(r'[| :\-]+',row)]
            t=doc.add_table(rows=0,cols=len(parsed[0]));t.style='Table Grid'
            for j,row in enumerate(parsed):
                cells=t.add_row().cells
                for c,value in zip(cells,row):
                    c.text=value
                    for p in c.paragraphs:
                        p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4)
                        for r in p.runs:r.font.size=Pt(9);r.bold=j==0
                trpr=t.rows[-1]._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
                if j==0:
                    trpr.append(OxmlElement('w:tblHeader'))
                    for c in cells:
                        shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E9F0EC');c._tc.get_or_add_tcPr().append(shade)
            doc.add_paragraph();continue
        if line.startswith('# '):doc.add_paragraph(line[2:].replace('的三维','的\n三维'),style='Title')
        elif line.startswith('## '):
            if line.startswith('## 三'):picture('architecture.png','图1 已实现的声音生成与数据归档流程')
            if line.startswith('## 四'):
                picture('workbench.png','图2 本机创作工作台实际界面')
                picture('timeline.png','图3 历史成品声音时间线与原文对应报告')
            doc.add_heading(line[3:],level=1)
        elif line.startswith('### '):doc.add_heading(line[4:],level=2)
        else:
            p=doc.add_paragraph(line)
            if not line.startswith(('[','项目','指导','报告','签名','学院')):p.paragraph_format.first_line_indent=Cm(.74)
            if line.startswith('['):
                for r in p.runs:r.font.size=Pt(9)
    doc.core_properties.author='项目组';doc.core_properties.title='三维动态音景自动生成系统中期检查报告'
    doc.save(FOLDER/'中期检查报告.docx')
    print('DOCX created')

if __name__=='__main__':main()
